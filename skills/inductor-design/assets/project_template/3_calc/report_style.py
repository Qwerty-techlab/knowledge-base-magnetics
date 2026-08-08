# -*- coding: utf-8 -*-
"""Единый стиль пояснительной записки по ГОСТ Р 2.105—2019."""
from __future__ import annotations

import os
from datetime import date
from typing import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


FONT = "Times New Roman"
SZ_MAIN = Pt(14)
SZ_TABLE = Pt(10)


def fmt(value, digits: int = 2) -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, int):
        return str(value)
    try:
        return f"{float(value):.{digits}f}".replace(".", ",")
    except (TypeError, ValueError):
        return str(value)


def _add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = SZ_MAIN
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(10)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SZ_MAIN
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    first = section.first_page_footer.paragraphs[0]
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.first_line_indent = Cm(0)
    run = first.add_run(str(date.today().year))
    run.font.name = FONT
    run.font.size = SZ_MAIN
    _add_page_field(section.footer.paragraphs[0])

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    return doc


def h1(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.name = FONT
    run.font.size = SZ_MAIN
    return paragraph


def h2(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = FONT
    run.font.size = SZ_MAIN
    return paragraph


def para(doc, text: str, indent: bool = True, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(1.25 if indent else 0)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = SZ_MAIN
    return paragraph


def formula(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = SZ_MAIN
    run.italic = True
    return paragraph


def table_caption(doc, number: int, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"Таблица {number} — {text}")
    run.font.name = FONT
    run.font.size = SZ_MAIN
    return paragraph


def fig_caption(doc, number: int, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"Рисунок {number} — {text}")
    run.font.name = FONT
    run.font.size = SZ_MAIN
    return paragraph


def _set_cell(cell, text: str, bold: bool = False,
              align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = SZ_TABLE


def add_table(doc, header: Sequence[str], rows: Iterable[Sequence], widths=None):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(header):
        _set_cell(table.rows[0].cells[index], str(value), bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            align = (WD_ALIGN_PARAGRAPH.LEFT if index == 0
                     else WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell(cells[index], str(value), align=align)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def add_figure(doc, path: str, caption: str, counter, width: float = 15.5):
    if not os.path.exists(path):
        raise FileNotFoundError(f"не найден обязательный рисунок отчёта: {path}")
    doc.add_picture(path, width=Cm(width))
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    fig_caption(doc, counter[0], caption)
    counter[0] += 1
    return paragraph
