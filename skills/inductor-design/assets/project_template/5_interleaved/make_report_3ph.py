# -*- coding: utf-8 -*-
"""Полный Word-отчёт для многофазного interleaved boost-дросселя.

Расчётные данные берутся только из актуального ``results_interleaved.json``.
Оформление и последовательность разделов задаются единым модулем
``3_calc/report_style.py`` и не зависят от резервных копий старых проектов.

Отчёт имеет предварительный аналитический статус: результаты FEM, измерения
L_diff(I), тепловые испытания, первый резонанс и ЭМС не объявляются выполненными.
"""
from __future__ import annotations

import importlib.util
import json
import math
import os
import re
from datetime import date
from typing import Any, Dict, Iterable, List, Tuple

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
HELPER_PATH = os.path.join(ROOT, "3_calc", "report_style.py")
JSON_PATH = os.path.join(HERE, "results_interleaved.json")
OUT_DIR = os.path.join(ROOT, "1_output_files")
IMG_DIR = os.path.join(OUT_DIR, "img_3ph")
OUT_PATH = os.path.join(OUT_DIR, "Расчёт_дросселя_interleaved.docx")
FKEYS: Tuple[Tuple[str, str], ...] = ()
REFS: Dict[str, int] = {}
FORMULA_NO = [1]


def _load_report_helpers():
    spec = importlib.util.spec_from_file_location("drossel_report_helpers", HELPER_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"не удалось загрузить функции оформления: {HELPER_PATH}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


mr = _load_report_helpers()
fmt = mr.fmt
h1, h2, para = mr.h1, mr.h2, mr.para
table_caption, add_table, add_figure = mr.table_caption, mr.add_table, mr.add_figure


def formula(doc, text: str):
    """Добавить формулу с единым сквозным номером у правого поля."""
    number = FORMULA_NO[0]
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(8.2), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.italic = True
    paragraph.add_run("\t")
    run = paragraph.add_run(f"({number})")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    FORMULA_NO[0] += 1
    return paragraph


def setup_document():
    """Создать документ в стиле Backup и исправить его OOXML-поля."""
    doc = mr.setup_document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend((begin, instruction, separate, end))
    return doc


def sel(res: Dict[str, Any], f_key: str) -> Dict[str, Any]:
    return res[f_key]["selected"]


def _frequency_label(frequency_hz: float) -> str:
    value = f"{frequency_hz / 1e3:.6f}".rstrip("0").rstrip(".")
    return f"{value.replace('.', ',')} кГц"


def _frequency_pairs(data: Dict[str, Any]) -> Tuple[Tuple[str, str], ...]:
    pairs = []
    for key, result in data.get("results", {}).items():
        frequency_hz = result.get("f_sw_Hz")
        if frequency_hz is None:
            frequency_hz = (
                result.get("selected", {}).get("mode", {}).get("f_sw_kHz")
            )
            if frequency_hz is not None:
                frequency_hz *= 1e3
        if frequency_hz is None:
            raise SystemExit(f"для варианта {key} не задана частота коммутации")
        pairs.append((key, _frequency_label(float(frequency_hz))))
    if not pairs:
        raise SystemExit("в results_interleaved.json отсутствуют частотные варианты")
    return tuple(pairs)


def _frequency_text() -> str:
    return ", ".join(label for _, label in FKEYS)


def _frequency_header(first: str = "Величина", last: str | None = "Ед.") -> List[str]:
    header = [first, *(label for _, label in FKEYS)]
    if last is not None:
        header.append(last)
    return header


def _frequency_widths(first: float, variants_total: float,
                      last: float | None = None) -> List[float]:
    width = variants_total / len(FKEYS)
    widths = [first, *(width for _ in FKEYS)]
    if last is not None:
        widths.append(last)
    return widths


def _system_loss(selected: Dict[str, Any], n_phases: int) -> float:
    losses = selected["losses"]
    if "P_total_system_W" in losses:
        return losses["P_total_system_W"]
    if "P_total_3ph_W" in losses and n_phases == 3:
        return losses["P_total_3ph_W"]
    return losses["P_total_W"] * n_phases


def _yes_no(value: Any) -> str:
    if value is True:
        return "да"
    if value is False:
        return "НЕТ"
    return "ожидается"


def _short_status(status: Any) -> str:
    mapping = {
        "analytical_pass": "аналитически допустим",
        "calculated": "рассчитано",
        "calculated_80_harmonics": "рассчитано, 80 гармоник",
        "preliminary": "предварительно",
        "preliminary_model": "предварительная модель",
        "preliminary_thermal_model": "предварительная тепловая модель",
        "analytical_preliminary": "аналитический предварительный",
        "preliminary_dipole_estimate": "предварительная дипольная оценка",
        "pending_fem_or_measurement": "ожидает FEM/измерения",
        "pending_converged_fem": "ожидает FEM со сходимостью",
        "pending_measurement": "ожидает измерения",
    }
    return mapping.get(status, str(status or "—"))


def _report_text(value: Any) -> str:
    """Привести диагностический текст расчёта к обозначениям отчёта."""
    text = str(value or "—")
    text = re.sub(r"(?<=\d)\.(?=\d)", ",", text)
    for source, target in (
        ("°C", "°С"), ("A/mm²", "А/мм²"), ("A/mm2", "А/мм²"),
        ("uH", "мкГн"), ("µH", "мкГн"), ("mT", "мТл"),
    ):
        text = text.replace(source, target)
    return text


def _captioned_table(doc, tn, caption, header, rows, widths=None):
    table_caption(doc, tn[0], caption)
    add_table(doc, header, rows, widths=widths)
    tn[0] += 1


def _frequency_values(res, path: Iterable[str], digits=2):
    values = []
    for key, _ in FKEYS:
        value: Any = sel(res, key)
        for part in path:
            value = value[part]
        values.append(fmt(value, digits))
    return values


def title_page(doc, tz):
    for _ in range(3):
        para(doc, "")
    p = para(
        doc,
        f"РАСЧЁТ ДРОССЕЛЯ ФАЗЫ {tz['n_phases']}-ФАЗНОГО ПОВЫШАЮЩЕГО "
        "(INTERLEAVED BOOST) ПРЕОБРАЗОВАТЕЛЯ",
        indent=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(16)
    para(doc, "")
    para(doc, "Пояснительная записка", indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        doc,
        f"Выходная мощность {fmt(tz['p_out_nom'], 0)} Вт; "
        f"{fmt(tz['u_in_nom'], 0)} В / {fmt(tz['u_out_nom'], 0)} В; "
        f"число фаз {fmt(tz['n_phases'])}; КПД {fmt(tz['eta'] * 100, 0)} %",
        indent=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    para(doc, "Аналитический предварительный расчёт", indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(8):
        para(doc, "")
    para(
        doc,
        "Расчёт выполнен по исправленной методике проектирования "
        "высокочастотных силовых дросселей "
        "(Metodika_VCh_silovykh_drosseley_bez_rascheta.tex). "
        "Оформление по ГОСТ Р 2.105.",
        indent=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for _ in range(4):
        para(doc, "")
    para(doc, str(date.today().year), indent=False,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sec_input(doc, data, tn):
    tz, der, meta = data["tz"], data["tz_derived"], data["meta"]
    h1(doc, "1 Исходные данные, режимы и допуски")
    para(
        doc,
        f"Расчёт выполняется для одного дросселя фазы. В {tz['n_phases']}-фазной "
        f"системе установлено {tz['n_phases']} одинаковых фазных дросселей; "
        f"варианты {_frequency_text()} являются альтернативными конструкциями.",
    )
    rows = [
        ["Число фаз", fmt(tz["n_phases"]), "—", "ТЗ"],
        ["Входное напряжение U_вх", fmt(tz["u_in_nom"], 0), "В", "ТЗ"],
        ["Выходное напряжение U_вых", fmt(tz["u_out_nom"], 0), "В", "ТЗ"],
        ["Выходной ток I_вых", fmt(tz["i_out_nom"], 0), "А", "ТЗ"],
        ["Выходная мощность P_вых", fmt(tz["p_out_nom"], 0), "Вт", "ТЗ"],
        ["КПД η", fmt(tz["eta"] * 100, 0), "%", "ТЗ"],
        ["Частоты коммутации",
         "; ".join(label.removesuffix(" кГц") for _, label in FKEYS),
         "кГц", "ТЗ"],
        ["Пульсация тока фазы r_I", fmt(tz["r_i"] * 100, 0), "%", "ТЗ"],
        ["Пульсации входного/выходного напряжения",
         fmt(tz["ripple_u_in"] * 100, 0), "%", "ТЗ"],
        ["Допуск входного напряжения", "±" + fmt(tz["tol_u_in"] * 100, 0),
         "%", "методика"],
        ["Допуск индуктивности", "±" + fmt(tz["tol_L"] * 100, 0),
         "%", "методика"],
        ["Перегрузка по току", "+" + fmt(tz["tol_i_load"] * 100, 0),
         "%", "методика"],
        ["Разброс токов фаз", "±" + fmt(tz["tol_i_share"] * 100, 0),
         "%", "методика"],
        ["Температура среды", fmt(tz["t_amb"], 0), "°С", "методика"],
        ["Предел температуры обмотки", fmt(tz["t_wind_max"], 0),
         "°С", "методика"],
        ["Предел температуры сердечника", fmt(tz["t_core_max"], 0),
         "°С", "методика"],
        ["Коэффициент B для аналитического отбора k_B,предв",
         fmt(tz["k_B_prelim"], 2), "—", "методика"],
        ["Коэффициент B для финальной FEM-проверки k_B",
         fmt(tz["k_B"], 2), "—", "методика"],
        ["Допустимое заполнение медью k_Cu", fmt(tz["k_cu_dop"], 2),
         "—", "методика"],
        ["Допустимая занятость окна k_зан", fmt(tz["k_zan_dop"], 2),
         "—", "методика"],
        ["Допустимая плотность тока J_доп", fmt(tz["j_dop"] / 1e6, 1),
         "А/мм²", "методика"],
        ["Число физических зазоров на дроссель", fmt(tz["n_gaps"]),
         "—", "ТЗ/методика"],
        ["Запас по первому резонансу k_f", fmt(tz["k_f_res"], 0),
         "—", "методика"],
        ["Требуемый запас по ЭМС", fmt(tz["M_EMC"], 0), "дБ", "методика"],
    ]
    REFS["input"] = tn[0]
    _captioned_table(doc, tn, "Исходные данные, допуски и их источники",
                     ["Величина", "Значение", "Ед.", "Источник"], rows,
                     [8.6, 3.2, 1.9, 3.1])
    h2(doc, "1.1 Производные исходные величины")
    formula(doc, "P_вх = P_вых / η;   I_вх = P_вх / U_вх;   I_ф = I_вх / n_ф")
    p_ui = tz["u_out_nom"] * tz["i_out_nom"]
    base_text = (f"Получено P_вх = {fmt(der['P_in_W'], 1)} Вт, "
                 f"I_вх = {fmt(der['I_in_A'], 2)} А, I_ф = "
                 f"{fmt(der['I_phase_dc_A'], 2)} А. ")
    if math.isclose(p_ui, tz["p_out_nom"], rel_tol=1e-9, abs_tol=1e-6):
        para(doc, base_text + f"Произведение U_вых·I_вых = {fmt(p_ui, 0)} Вт "
             "совпадает с заданной выходной мощностью.")
    else:
        para(doc, base_text + f"Произведение U_вых·I_вых равно {fmt(p_ui, 0)} Вт "
             f"и не совпадает с P_вых = {fmt(tz['p_out_nom'], 0)} Вт; "
             f"консервативно главенствует заданная мощность "
             f"{fmt(tz['p_out_nom'] / 1e3, 3)} кВт.")
    para(doc, f"Статус исходного расчёта: {_short_status(meta['status'])}; финальная "
              "верификация не выполнена.")


def sec_mode(doc, data, tn, fn):
    res = data["results"]
    h1(doc, "2 Электрический режим дросселя фазы")
    h2(doc, "2.1 Скважность, токи и индуктивность")
    para(doc, "Для идеализированной boost-модели падения напряжения на ключе, "
              "диоде и сопротивлении обмотки в балансе вольт-секунд не учитываются.")
    formula(doc, "D = 1 − U_вх/U_вых;   L_треб = U_вх·D/(f·ΔI_пп)")
    rows = []
    labels = [
        ("Скважность D", "D", "—", 4),
        ("Ошибка вольт-секундного баланса", "volt_second_error_V", "В", 6),
        ("Постоянный ток фазы I_DC", "I_dc_A", "А", 3),
        ("Размах пульсации фазы ΔI_пп", "dI_pp_A", "А", 3),
        ("Максимальный ток фазы", "I_max_A", "А", 3),
        ("Действующий ток фазы", "I_rms_A", "А", 3),
        ("Требуемая индуктивность", "L_req_uH", "мкГн", 3),
        ("Целевая номинальная индуктивность", "L_nom_target_uH", "мкГн", 3),
        ("Максимальное потокосцепление", "psi_max_mWb", "мВб", 4),
        ("Максимальная энергия дросселя", "W_max_mJ", "мДж", 3),
    ]
    for name, field, unit, digits in labels:
        values = []
        for key, _ in FKEYS:
            source = res[key]["mode_converter"] if field in (
                "D", "volt_second_error_V") else res[key]["mode_phase"]
            values.append(fmt(source[field], digits))
        rows.append([name, *values, unit])
    _captioned_table(doc, tn, "Электрический режим одного фазного дросселя",
                     _frequency_header(), rows,
                     _frequency_widths(7.8, 6.2, 2.4))

    n_phases = data["tz"]["n_phases"]
    phase_shift = 360.0 / n_phases
    h2(doc, f"2.2 Компенсация пульсаций {n_phases}-фазным чередованием")
    rows = []
    for key, label in FKEYS:
        mode = res[key]["mode_converter"]
        rows.append([
            label, fmt(mode["I_in_dc_A"], 3), fmt(mode["K_ripple"], 4),
            fmt(mode["K_ripple_numeric"], 4), fmt(mode["dI_in_pp_A"], 4),
            fmt(mode["dI_in_rel_pct"], 3), fmt(mode["f_in_ripple_Hz"] / 1e3, 0),
        ])
    _captioned_table(doc, tn, "Пульсация суммарного входного тока",
                     ["f", "I_вх, А", "K расч.", "K числ.", "ΔI_вх, А",
                      "ΔI_вх, %", "f_пульс, кГц"], rows,
                     [2.2, 2.4, 2.1, 2.1, 2.5, 2.5, 2.5])
    para(doc, f"Коэффициент подавления получен суммированием {n_phases} токов, "
              f"сдвинутых на {fmt(phase_shift, 3)}°. Совпадение аналитического и численного "
              "значений служит внутренней проверкой расчёта.")
    add_figure(doc, os.path.join(IMG_DIR, "ripple_vs_D.png"),
               f"Коэффициент подавления пульсаций {n_phases}-фазной схемой", fn)
    for key, label in FKEYS:
        add_figure(doc, os.path.join(IMG_DIR, f"waveform_{key}.png"),
                   f"Токи фаз, суммарный входной ток и напряжение на дросселе, {label}", fn)


def _best_summary(best):
    if not best:
        return ["—", "—", "—", "—"]
    return [fmt(best["losses"]["P_total_W"], 3),
            fmt(best["design"]["N"], 0),
            fmt(best["design"]["B_max_wc_mT"], 1),
            fmt(best["design"]["gap_mm"], 3)]


def sec_core(doc, data, tn):
    res, materials, tz = data["results"], data["materials"], data["tz"]
    h1(doc, "3 Выбор магнитопровода, материала и провода")
    preferred = ", ".join(tz.get("preferred_core_families", ()))
    para(doc, "Выбор выполнен перебором сердечников, доступных сочетаний "
              "материалов, литцендрата, числа параллельных пучков и числа витков. "
              f"Семейства {preferred} имеют приоритет; "
              "прямоугольное семейство E допускается только при отсутствии "
              "аналитически пригодного предпочтительного варианта.")

    rows = []
    for key, label in FKEYS:
        e = res[key]["enumeration"]
        rows.append([label, e["n_seen"], e["n_calculated"], e["n_feasible"],
                     e["n_pareto"], e["n_in_window"],
                     fmt(e["loss_window_pct"], 1), _short_status(e["selection_status"])])
    _captioned_table(doc, tn, "Объём перебора и отбор решений",
                     ["f", "Просм.", "Рассч.", "Допуст.", "Парето", "В окне",
                      "Окно, %", "Статус"], rows,
                     [1.8, 2.0, 2.0, 2.0, 1.8, 1.8, 2.0, 3.2])

    for key, label in FKEYS:
        rows = []
        for item in res[key]["area_product_screening"]:
            if not item.get("core"):
                continue
            rows.append([
                item["core"], fmt(item["A_e_mm2"], 0), fmt(item["A_w_mm2"], 0),
                fmt(item["Ap_catalog_mm4"], 0), fmt(item["Ap_required_mm4"], 0),
                _yes_no(item["ok"]), _yes_no(item["custom_gap_possible"]),
                item["catalog_gap_count"],
            ])
        _captioned_table(doc, tn, f"Предварительный отбор по произведению площадей, {label}",
                         ["Сердечник", "A_e, мм²", "A_w, мм²", "A_p, мм⁴",
                          "A_p,треб, мм⁴", "A_p", "Заказной g", "Кат. g"], rows,
                         [3.2, 2.0, 2.0, 2.5, 2.8, 1.4, 2.0, 1.6])

        rows = []
        for item in res[key].get("turn_window_screening", []):
            if not item.get("core"):
                continue
            rows.append([
                item["core"], item.get("family", "—"),
                fmt(item["A_min_mm2"], 1), fmt(item["A_N_mm2"], 1),
                item["N_by_induction"], item["N_by_window"],
                _yes_no(item["ok"]), _report_text(item["verdict"]), item.get("source", "—"),
            ])
        _captioned_table(doc, tn, f"Проверка требуемых витков и окна, {label}",
                         ["Типоразмер", "Сем.", "A_min, мм²", "A_N, мм²",
                          "N_B", "N_w", "Итог", "Численная причина", "Источник"], rows,
                         [2.7, 1.2, 1.7, 1.7, 1.1, 1.1, 1.2, 3.9, 3.8])

    h2(doc, "3.1 Полное сравнение сердечников и материалов")
    for key, label in FKEYS:
        rows = []
        chosen = sel(res, key)
        for item in res[key]["core_material_table"]:
            best = item.get("best")
            p, n, b_wc, gap = _best_summary(best)
            selected_mark = "выбран" if (
                item["core"] == chosen["core"]["name"] and
                item["material"] == chosen["material"]["name"]
            ) else "—"
            rows.append([
                item["core"], item["material"], item["n_feasible"],
                p, n, b_wc, gap, selected_mark,
                _report_text(item.get("reason") or _short_status(item["status"])),
            ])
        _captioned_table(doc, tn, f"Сочетания магнитопровод × материал, {label}",
                         ["Сердечник", "Мат.", "Годных", "P, Вт", "N",
                          "B_wc, мТл", "g, мм", "Выбор", "Статус"], rows,
                         [3.0, 1.5, 1.7, 1.8, 1.2, 2.2, 1.8, 1.7, 3.4])

    rows = []
    for key, mat in materials.items():
        st = mat["steinmetz"]
        rows.append([
            key, fmt(mat["mu_i"], 0), fmt(mat["B_S_25C_mT"], 0),
            fmt(mat["B_S_100C_mT"], 0),
            f"{fmt(mat['f_min_kHz'], 0)}…{fmt(mat['f_max_kHz'], 0)}",
            fmt(mat["T_curie_C"], 0), fmt(st["k"], 5),
            fmt(st["alpha"], 4), fmt(st["beta"], 3),
        ])
    _captioned_table(doc, tn, "Паспортные и аппроксимационные параметры материалов",
                     ["Марка", "μ_i", "B_S25, мТл", "B_S100, мТл", "f, кГц",
                      "T_C, °С", "k", "α", "β"], rows,
                     [1.7, 1.5, 2.1, 2.2, 2.4, 1.8, 2.1, 1.7, 1.7])
    para(doc, "Параметры Штейнмеца имеют ограничения идентификации, указанные "
              "в базе данных; поэтому потери в феррите имеют предварительный статус.")

    for key, label in FKEYS:
        rows = []
        for mat_name, item in res[key]["material_comparison"].items():
            rows.append([
                mat_name, _yes_no(item.get("in_range")),
                fmt(item.get("B_allow_mT"), 1), fmt(item.get("B_max_wc_mT"), 1),
                fmt(item.get("P_fe_W"), 4),
                fmt(item.get("P_total_W", item.get("P_W")), 3),
                _yes_no(item.get("analytical_ok")),
                "числовая оценка" if item.get("P_total_W", item.get("P_W")) is not None
                else item.get("status", "нет паспортного A_L0 для выбранного сердечника"),
            ])
        chosen_name = sel(res, key)["core"]["name"]
        _captioned_table(doc, tn, f"Материалы на выбранном {chosen_name}, {label}",
                         ["Марка", "В диапазоне", "B_доп, мТл", "B_wc, мТл",
                          "P_Fe, Вт", "P_Σ, Вт", "Допуст.", "Примечание"], rows,
                         [1.6, 2.0, 2.2, 2.2, 1.9, 1.9, 1.7, 4.0])

    h2(doc, "3.2 Выбранные компоненты и обоснование")
    rows = []
    for key, label in FKEYS:
        b = sel(res, key)
        rows.append([
            label, b["core"]["name"], b["material"]["name"],
            b["core"]["part_ungapped"], b["wire"]["designation"],
            b["design"]["N"], fmt(b["losses"]["P_total_W"], 3),
        ])
    _captioned_table(doc, tn, "Выбранные магнитопровод, материал и провод",
                     ["f", "Сердечник", "Мат.", "Код половинки", "Провод",
                      "N", "P_1ф, Вт"], rows,
                     [1.9, 3.0, 1.5, 3.3, 4.4, 1.2, 2.1])
    for key, label in FKEYS:
        chosen = sel(res, key)
        alternatives = [x["best"] for x in res[key]["core_material_table"]
                        if x.get("best") and not (
                            x["core"] == chosen["core"]["name"] and
                            x["material"] == chosen["material"]["name"])]
        if alternatives:
            nearest = min(alternatives, key=lambda x: x["losses"]["P_total_W"])
            delta = (chosen["losses"]["P_total_W"] /
                     nearest["losses"]["P_total_W"] - 1.0) * 100.0
            relation = (f"на {fmt(abs(delta), 1)} % больше" if delta >= 0.0
                        else f"на {fmt(abs(delta), 1)} % меньше")
            policy_note = (" Семейство E не выбрано, поскольку в предпочтительных "
                           "семействах есть допустимые варианты."
                           if res[key]["enumeration"].get("preferred_feasible", 0) else
                           " Семейство E применено как резервное после отказа предпочтительных вариантов.")
            para(doc, f"Для {label} выбран {chosen['core']['name']}/"
                      f"{chosen['material']['name']}: "
                      f"{fmt(chosen['losses']['P_total_W'], 3)} Вт на фазу. "
                      f"Наименьшие потери среди остальных допустимых сочетаний имеет "
                      f"{nearest['core']['name']}/{nearest['material']['name']}: "
                      f"{fmt(nearest['losses']['P_total_W'], 3)} Вт; выбранный вариант "
                      f"даёт потери {relation}.{policy_note}")


def sec_magnetic(doc, data, tn):
    res = data["results"]
    q = res[FKEYS[0][0]]["quantities"]
    n_gaps = data["tz"]["n_gaps"]
    h1(doc, "4 Расчёт числа витков, зазора, индукции и энергии")
    para(doc, f"Для одного дросселя принято комплектов магнитопровода — "
              f"{q['core_sets_per_inductor']}, половинок — "
              f"{q['core_halves_per_inductor']}. Число физических зазоров — "
              f"{n_gaps}; распределённый зазор не применяется.")
    formula(doc, "A_L = L/N²;   B = Ψ/(N·A_min);   W = L·I²/2")
    labels = [
        ("Число витков N", "N", "—", 0),
        ("Требуемая индуктивность L_треб", None, "мкГн", 3),
        ("Номинальная индуктивность L", "L_nom_uH", "мкГн", 3),
        ("Минимальная индуктивность L_min", "L_min_uH", "мкГн", 3),
        ("Требуемое A_L", "A_L_req_nH", "нГн/вит²", 2),
        ("Расчётное A_L", "A_L_actual_nH", "нГн/вит²", 2),
        ("Один физический зазор g", "gap_mm", "мм", 3),
        ("Минимальная оценка g", "gap_model_min_mm", "мм", 3),
        ("Максимальная оценка g", "gap_model_max_mm", "мм", 3),
        ("Коэффициент выпучивания F_f", "F_fringe", "—", 4),
        ("Индукция B_wc", "B_max_wc_mT", "мТл", 2),
        ("Предварительный предел B", "B_allow_mT", "мТл", 2),
        ("Размах индукции ΔB", "dB_pp_mT", "мТл", 3),
        ("Отношение сопротивлений k_R", "k_R", "—", 2),
    ]
    rows = []
    for name, field, unit, digits in labels:
        values = []
        for key, _ in FKEYS:
            b = sel(res, key)
            value = b["mode"]["L_req_uH"] if field is None else b["design"][field]
            values.append(fmt(value, digits))
        rows.append([name, *values, unit])
    _captioned_table(doc, tn, "Магнитный расчёт выбранных вариантов",
                     _frequency_header(), rows,
                     _frequency_widths(7.5, 6.2, 2.6))
    for key, label in FKEYS:
        b = sel(res, key)
        para(doc, f"Для {label} зазор задаётся целевым A_L = "
                  f"{fmt(b['design']['A_L_req_nH'], 2)} нГн/вит². Размер "
                  f"g = {fmt(b['design']['gap_mm'], 3)} мм является аналитической "
                  "оценкой и должен быть уточнён по FEM и измерению A_L после обработки.")
    core_names = sorted({sel(res, key)["core"]["name"] for key, _ in FKEYS})
    part_numbers = sorted({sel(res, key)["core"].get("part_ungapped") or "по заказу"
                           for key, _ in FKEYS})
    para(doc, "Готовые каталожные исполнения с требуемым A_L для выбранных "
              f"магнитопроводов ({', '.join(core_names)}) отсутствуют; используются "
              f"исходные половинки {', '.join(part_numbers)} с последующей обработкой "
              "центрального стержня по целевому A_L.")


def sec_winding(doc, data, tn, fn):
    res = data["results"]
    h1(doc, "5 Укладка обмотки и расчёт заполнения окна")
    formula(doc, "k_Cu = N·S_Cu/A_N,эф;   k_зан = N·S_наруж/A_N,эф")
    labels = [
        ("Параллельных пучков", ("wire", "n_parallel"), "—", 0),
        ("Суммарное сечение меди", ("wire", "S_cu_mm2"), "мм²", 3),
        ("Число слоёв", ("design", "n_layers"), "—", 0),
        ("Высота укладки", ("design", "h_wind_mm"), "мм", 2),
        ("Радиальная ширина", ("design", "b_winding_mm"), "мм", 2),
        ("Средняя длина витка", ("design", "l_turn_mm"), "мм", 2),
        ("Длина одного пучка", ("design", "wire_length_m"), "м", 2),
        ("Общая длина на один дроссель", ("design", "wire_length_total_m"), "м", 2),
        ("Общая длина на систему", ("design", "wire_length_total_system_m"), "м", 2),
        ("Заполнение медью k_Cu", ("design", "k_cu"), "—", 4),
        ("Геометрическая занятость k_зан", ("design", "k_zan"), "—", 4),
    ]
    rows = []
    for name, path, unit, digits in labels:
        values = []
        for key, _ in FKEYS:
            b = sel(res, key)
            values.append(fmt(b[path[0]][path[1]], digits))
        rows.append([name, *values, unit])
    _captioned_table(doc, tn, "Укладка обмотки и заполнение окна",
                     _frequency_header(), rows,
                     _frequency_widths(7.5, 6.2, 2.6))
    para(doc, "Каркас, изоляция, отступ обмотки от зазора и крепление должны "
              "быть зафиксированы рабочим чертежом и проверены по фактической "
              "геометрии выбранного магнитопровода.")
    for key, label in FKEYS:
        add_figure(doc, os.path.join(IMG_DIR, f"winding_{key}.png"),
                   f"Расчётная укладка обмотки, {label}", fn)


def sec_losses(doc, data, tn, fn):
    res, n_phases = data["results"], data["tz"]["n_phases"]
    h1(doc, "6 Потери по составляющим и температурный расчёт")
    para(doc, "Потери рассчитаны раздельно для постоянной и переменной "
              "составляющих тока. Потери меди в локальном поле зазора не "
              "включены и требуют FEM.")
    labels = [
        ("Сопротивление горячей обмотки", "R_dc_Ohm", "Ом", 6),
        ("Потери постоянного тока P_DC", "P_dc_W", "Вт", 4),
        ("Потери переменного тока P_AC", "P_ac_W", "Вт", 5),
        ("Потери в меди P_Cu", "P_cu_W", "Вт", 4),
        ("Потери в феррите P_Fe", "P_fe_W", "Вт", 5),
        ("Суммарные потери одного дросселя", "P_total_W", "Вт", 4),
        (f"Суммарные потери {n_phases} дросселей", None, "Вт", 4),
        ("Доля потерь от мощности системы", "loss_pct_system", "%", 4),
        ("Температура обмотки", "t_winding_C", "°С", 2),
        ("Температура сердечника", "t_core_C", "°С", 2),
        ("Перегрев", "dtheta_K", "К", 2),
    ]
    rows = []
    for name, field, unit, digits in labels:
        values = [
            fmt(_system_loss(sel(res, key), n_phases), digits)
            if field is None else fmt(sel(res, key)["losses"][field], digits)
            for key, _ in FKEYS
        ]
        rows.append([name, *values, unit])
    _captioned_table(doc, tn, "Потери и предварительный температурный расчёт",
                     _frequency_header(), rows,
                     _frequency_widths(7.5, 6.2, 2.6))
    for key, label in FKEYS:
        rows = [[fmt(h["f_Hz"] / 1e3, 0), fmt(h["I_rms_A"], 5),
                 fmt(h["F_R"], 3), fmt(h["R_ac_Ohm"], 6), fmt(h["P_W"], 6)]
                for h in sel(res, key)["harmonics"][:8]]
        _captioned_table(doc, tn, f"Потери переменного тока по гармоникам, {label}",
                         ["f, кГц", "I_скз, А", "F_R", "R_ac, Ом", "P, Вт"], rows,
                         [2.7, 3.0, 2.5, 3.4, 3.4])
    para(doc, "Тепловая модель сосредоточенная и предварительная. Температурное "
              "поле и локальные горячие точки должны быть проверены FEM и прототипом.")
    for key, label in FKEYS:
        add_figure(doc, os.path.join(IMG_DIR, f"phases_{key}.png"),
                   f"Фазные токи для расчёта потерь, {label}", fn)


def sec_parasitics(doc, data, tn, fn):
    res = data["results"]
    h1(doc, "7 Собственная ёмкость, первый резонанс и импеданс")
    formula(doc, "f_рез = 1/(2π·√(L·C_соб))")
    rows = []
    labels = [
        ("Собственная ёмкость C_соб", "C_self_pF", "пФ", 2),
        ("Предельно допустимая ёмкость", "C_self_max_pF", "пФ", 2),
        ("Первый собственный резонанс", "f_res_MHz", "МГц", 3),
    ]
    for name, field, unit, digits in labels:
        rows.append([name, *[fmt(sel(res, key)["design"][field], digits)
                             for key, _ in FKEYS], unit])
    _captioned_table(doc, tn, "Оценка паразитных параметров",
                     _frequency_header(), rows,
                     _frequency_widths(7.5, 6.2, 2.6))
    para(doc, "Расчётная частота резонанса превышает требуемую границу, однако "
              "окончательное значение определяется измерением готового дросселя.")
    add_figure(doc, os.path.join(IMG_DIR, "impedance.png"),
               "Расчётная частотная зависимость модуля импеданса", fn)


def sec_emc(doc, data, tn):
    tz, res = data["tz"], data["results"]
    h1(doc, "8 Электромагнитная совместимость")
    para(doc, "Вносимое ослабление и поле рассеяния приведены как результаты "
              "аналитической модели. Соответствие норме ЭМС не подтверждено "
              "до измерений на преобразователе.")
    first_key = FKEYS[0][0]
    keys = tuple(sel(res, first_key)["emc"]["A_Z_model_dB"])
    rows = []
    for freq in keys:
        rows.append([
            str(freq).replace("kHz", ""),
            *(fmt(sel(res, key)["emc"]["A_Z_model_dB"][freq], 2)
              for key, _ in FKEYS),
        ])
    _captioned_table(doc, tn, "Модель вносимого ослабления в полосе кондуктивных помех",
                     ["f, кГц", *(f"{label}, дБ" for _, label in FKEYS)], rows,
                     _frequency_widths(5.0, 11.2))
    rows = []
    for key, label in FKEYS:
        s = res[key]["stray_field"]
        rows.append([
            label, fmt(s["r_probe_mm"], 0), fmt(s["B_stray_pk_uT"], 4),
            fmt(s["B_stray_rms_uT"], 4), fmt(s["U_induced_rms_mV"], 4),
            _short_status(s["status"]),
        ])
    _captioned_table(doc, tn, "Оценка поля рассеяния зазора",
                     ["f", "r, мм", "B_pk, мкТл", "B_скз, мкТл",
                      "U_нав, мВ", "Статус"], rows,
                     [2.2, 2.2, 3.0, 3.0, 3.0, 3.8])
    para(doc, f"Финальный запас не менее {fmt(tz['M_EMC'], 0)} дБ должен быть "
              "подтверждён измерением во всей нормируемой полосе. Взаимная "
              f"индуктивность {tz['n_phases']} дросселей также проверяется "
              "на общей компоновке.")


def _margin_display(margin):
    value, limit, unit = margin.get("value"), margin.get("limit"), margin.get("unit", "—")
    scale, out_unit, digits = 1.0, unit, 3
    if unit == "Гн":
        scale, out_unit, digits = 1e6, "мкГн", 3
    elif unit == "Тл":
        scale, out_unit, digits = 1e3, "мТл", 2
    elif unit == "м":
        scale, out_unit, digits = 1e3, "мм", 2
    elif unit == "Гц":
        scale, out_unit, digits = 1e-6, "МГц", 3
    return fmt(None if value is None else value * scale, digits), \
        fmt(None if limit is None else limit * scale, digits), out_unit


def sec_margins(doc, data, tn):
    res = data["results"]
    h1(doc, "9 Запасы по обязательным ограничениям")
    para(doc, "Аналитические критерии и финальные проверки показаны раздельно. "
              "Пустое числовое значение означает, что требуется FEM или измерение.")
    for key, label in FKEYS:
        rows = []
        margins = sel(res, key)["margins"]
        for name, margin in margins.items():
            if name.startswith("_"):
                continue
            value, limit, unit = _margin_display(margin)
            rows.append([
                margin["name"], value, margin["cmp"], limit, unit,
                fmt(margin.get("margin_pct"), 2), _yes_no(margin.get("ok")),
                _short_status(margin.get("status")),
            ])
        _captioned_table(doc, tn, f"Запасы и незавершённые проверки, {label}",
                         ["Критерий", "Знач.", "Усл.", "Предел", "Ед.",
                          "Запас, %", "Вып.", "Статус"], rows,
                         [4.8, 2.0, 1.2, 2.0, 1.5, 1.8, 1.4, 3.5])
        ver = sel(res, key)["verification"]
        para(doc, f"Для {label}: аналитический отбор — "
                  f"{_yes_no(ver['analytical_pass'])}; финальная верификация — "
                  f"{_yes_no(ver['final_verified'])}.")


def sec_choice(doc, data, tn):
    res, n_phases = data["results"], data["tz"]["n_phases"]
    variants = [sel(res, key) for key, _ in FKEYS]
    h1(doc, "10 Сравнение частотных вариантов и рекомендация")
    rows = [
        ["Сердечник", *(b["core"]["name"] for b in variants)],
        ["Материал", *(b["material"]["name"] for b in variants)],
        ["Число витков", *(b["design"]["N"] for b in variants)],
        ["Параллельных пучков", *(b["wire"]["n_parallel"] for b in variants)],
        ["L_ном, мкГн", *(fmt(b["design"]["L_nom_uH"], 3) for b in variants)],
        ["Один зазор g, мм", *(fmt(b["design"]["gap_mm"], 3) for b in variants)],
        ["B_wc, мТл", *(fmt(b["design"]["B_max_wc_mT"], 2) for b in variants)],
        ["Потери одного дросселя, Вт",
         *(fmt(b["losses"]["P_total_W"], 4) for b in variants)],
        [f"Потери {n_phases} дросселей, Вт",
         *(fmt(_system_loss(b, n_phases), 4) for b in variants)],
        ["Температура обмотки, °С",
         *(fmt(b["losses"]["t_winding_C"], 2) for b in variants)],
        ["Первый резонанс, МГц",
         *(fmt(b["design"]["f_res_MHz"], 3) for b in variants)],
        ["Финально подтверждён",
         *(_yes_no(b["verification"]["final_verified"]) for b in variants)],
    ]
    _captioned_table(doc, tn, "Сравнение альтернативных конструкций",
                     _frequency_header("Показатель", None), rows,
                     _frequency_widths(8.0, 8.4))
    ranked = sorted(
        ((label, _system_loss(sel(res, key), n_phases)) for key, label in FKEYS),
        key=lambda item: item[1],
    )
    best_label, best_loss = ranked[0]
    if len(ranked) > 1:
        worst_label, worst_loss = ranked[-1]
        para(doc, f"Минимальные расчётные потери системы получены для варианта "
                  f"{best_label}: {fmt(best_loss, 3)} Вт. По сравнению с вариантом "
                  f"{worst_label} снижение составляет "
                  f"{fmt(worst_loss - best_loss, 3)} Вт.")
    else:
        para(doc, f"Для единственного рассчитанного варианта {best_label} потери "
                  f"системы составляют {fmt(best_loss, 3)} Вт.")
    para(doc, f"По потерям дросселей предварительно рекомендуется {best_label}. "
              "Это не финальный выбор: частота должна быть подтверждена измерением "
              "потерь всего преобразователя и электромагнитной совместимости.")


def sec_summary(doc, data, tn):
    res, n_phases = data["results"], data["tz"]["n_phases"]
    h1(doc, "11 Итоговые данные конструкции")
    para(doc, "Спецификации ниже относятся к взаимоисключающим вариантам "
              f"{_frequency_text()}. Количества приведены на {n_phases}-фазную систему.")
    for key, label in FKEYS:
        b, q = sel(res, key), res[key]["quantities"]
        rows = [
            ["Фазные дроссели", "одинаковая конструкция", q["phase_inductors"], "шт."],
            [f"Половинка {b['core']['name']}, {b['material']['name']}",
             b["core"]["part_ungapped"],
             q["core_halves_total"], "шт."],
            [f"Комплект магнитопровода из {q['core_halves_per_inductor']} половинок",
             "сборочная единица",
             q["core_sets_total"], "компл."],
            ["Обработка центрального стержня", f"целевое A_L = "
             f"{fmt(b['design']['A_L_req_nH'], 2)} нГн/вит²; зазоров: "
             f"{data['tz']['n_gaps']}",
             q["core_sets_total"], "компл."],
            ["Специальный каркас", "по рабочему чертежу",
             q["phase_inductors"], "шт."],
            ["Литцендрат", b["wire"]["designation"],
             fmt(b["design"]["wire_length_total_system_m"], 1), "м"],
            ["Число витков каждого дросселя", "—", b["design"]["N"], "вит."],
            ["Номинальная индуктивность", "—", fmt(b["design"]["L_nom_uH"], 2), "мкГн"],
            ["Расчётный зазор", "уточнить по FEM и измерению A_L",
             fmt(b["design"]["gap_mm"], 3), "мм"],
            ["Изоляционные и крепёжные материалы", "по рабочему чертежу",
             q["phase_inductors"], "компл."],
        ]
        _captioned_table(doc, tn, f"Предварительная спецификация, вариант {label}",
                         ["Позиция", "Обозначение/требование", "Количество", "Ед."],
                         rows, [5.2, 7.2, 2.3, 1.7])
        para(doc, f"Для {label}: один дроссель содержит "
                  f"{q['core_sets_per_inductor']} комплектов и "
                  f"{q['core_halves_per_inductor']} половинок; система содержит "
                  f"{q['core_sets_total']} комплектов, то есть "
                  f"{q['core_halves_total']} половинок.")


def sec_conclusions(doc, data):
    res, tz = data["results"], data["tz"]
    n_phases = tz["n_phases"]
    h1(doc, "12 Выводы")
    conclusions = [
        f"Выполнен предварительный аналитический расчёт {n_phases} одинаковых "
        f"фазных дросселей {n_phases}-фазного interleaved boost-преобразователя "
        f"мощностью {fmt(tz['p_out_nom'] / 1e3, 3)} кВт.",
        "Аналитические ограничения выполнены, но финальная верификация не завершена. "
        "Требуются FEM магнитного поля и потерь у зазора, проверка L_diff(I), "
        "тепловой прототип, измерение резонанса и ЭМС.",
    ]
    for key, label in FKEYS:
        b = sel(res, key)
        conclusions.insert(-1, (
            f"Вариант {label}: выбран {b['core']['name']} из {b['material']['name']}, "
            f"код половинки {b['core']['part_ungapped']}; N = {b['design']['N']}, "
            f"L = {fmt(b['design']['L_nom_uH'], 2)} мкГн, g = "
            f"{fmt(b['design']['gap_mm'], 3)} мм, потери одного дросселя "
            f"{fmt(b['losses']['P_total_W'], 3)} Вт, системы — "
            f"{fmt(_system_loss(b, n_phases), 3)} Вт."
        ))
    quantities = res[FKEYS[0][0]]["quantities"]
    conclusions.insert(-1, (
        f"Каждый дроссель содержит комплектов — "
        f"{quantities['core_sets_per_inductor']}, половинок — "
        f"{quantities['core_halves_per_inductor']}, физических зазоров — "
        f"{tz['n_gaps']}. На систему требуется "
        f"{quantities['core_sets_total']} комплектов, или "
        f"{quantities['core_halves_total']} половинок. Варианты "
        f"{_frequency_text()} нельзя суммировать в одной спецификации."
    ))
    best_key, best_label = min(
        FKEYS, key=lambda item: _system_loss(sel(res, item[0]), n_phases)
    )
    conclusions.append(
        f"По расчётным потерям дросселей предварительная системная рекомендация — "
        f"{best_label}; окончательная частота выбирается после измерения потерь "
        "всего преобразователя и помех."
    )
    for index, text in enumerate(conclusions, 1):
        p = para(doc, f"{index}) {text}")
        p.paragraph_format.space_after = Pt(6)


def sec_scope(doc, data, tn):
    res = data["results"]
    h1(doc, "13 Состав отчёта и вынесенное за пределы расчёта")
    rows = [
        ["1 Исходные данные, режимы и допуски", "разд. 1", "выполнено"],
        ["2 Источники паспортных величин", "разд. 1, 3", "выполнено"],
        ["3 Формы токов и гармонический состав", "разд. 2, 6", "выполнено"],
        ["4 Витки, зазор, индукция и энергия", "разд. 4", "предварительно"],
        ["5 Укладка и заполнение окна", "разд. 5", "предварительно"],
        ["6 Потери по компонентам и гармоникам", "разд. 6", "предварительно"],
        ["7 FEM со сходимостью сетки", "—", "не выполнено"],
        ["8 Температурное поле", "разд. 6", "сосредоточенная модель"],
        ["9 Импеданс и первый резонанс", "разд. 7", "модель; измерить"],
        ["10 ЭМС и внешнее поле", "разд. 8", "модель; измерить"],
        ["11 Запасы по ограничениям", "разд. 9", "аналитически"],
        ["12 Программа и протокол испытаний", "—", "после изготовления"],
    ]
    _captioned_table(doc, tn, "Соответствие обязательному составу отчёта",
                     ["Обязательный пункт", "Раздел", "Состояние"], rows,
                     [9.2, 3.0, 4.2])
    pending = sorted(set(item for key, _ in FKEYS
                         for item in sel(res, key)["verification"]["pending"]))
    translations = {
        "tune custom gap by measured A_L": "доводка зазора по измеренному A_L",
        "L_diff(I)": "измерение или FEM дифференциальной индуктивности L_diff(I)",
        "FEM B with physical edge radius": "FEM локальной B с физическим радиусом кромки",
        "FEM copper loss near gap": "FEM добавочных потерь меди у зазора",
        "thermal prototype": "тепловые испытания прототипа",
        "self-resonance measurement": "измерение первого собственного резонанса",
        "EMC measurement": "измерение ЭМС во всей нормируемой полосе",
    }
    para(doc, "До присвоения финального статуса необходимо выполнить:")
    for index, item in enumerate(pending, 1):
        para(doc, f"{index}) {translations.get(item, item)}")


def sec_sources(doc, data):
    h1(doc, "Список использованных источников")
    sources = [
        "Методика проектирования высокочастотных силовых дросселей "
        "Metodika_VCh_silovykh_drosseley_bez_rascheta.tex.",
        "ГОСТ Р 2.105—2019. Единая система конструкторской документации. "
        "Общие требования к текстовым документам.",
        "ГОСТ 8.417—2002. Государственная система обеспечения единства "
        "измерений. Единицы величин.",
        "ГОСТ 30805.22—2013 (CISPR 22:2006). Радиопомехи индустриальные. "
        "Нормы и методы измерений.",
        "Texas Instruments. Basic Calculation of a Boost Converter's Power Stage, SLVA372D.",
        "Texas Instruments. Multiphase Boost Converter Design and Interleaving Guidelines.",
        "Sullivan C. R., Zhang R. Y. Simplified Design Method for Litz Wire, APEC 2014.",
        "Massarini A., Kazimierczuk M. K. Self-Capacitance of Inductors, IEEE TPEL, 1997.",
        "Venkatachalam K. et al. Accurate Prediction of Ferrite Core Loss with Nonsinusoidal Waveforms, 2002.",
        "База компонентов knowledge_base/components_db: cores.json, ferrite_materials.json, litz_wire.json.",
    ]
    for key, label in FKEYS:
        selected = sel(data["results"], key)
        sources.extend([
            f"{selected['core']['datasheet']} — выбранный магнитопровод, {label}.",
            f"{selected['material']['datasheet']} — выбранный материал, {label}.",
            f"{selected['wire']['url']} — выбранный литцендрат, {label}.",
        ])
    for material in data.get("materials", {}).values():
        sources.append(f"{material['datasheet']} — материал сравнительного расчёта.")
    sources = list(dict.fromkeys(source for source in sources if source.strip()))
    for index, source in enumerate(sources, 1):
        p = para(doc, f"{index}. {source}")
        p.paragraph_format.space_after = Pt(6)


def main():
    global FKEYS
    FORMULA_NO[0] = 1
    with open(JSON_PATH, encoding="utf-8") as stream:
        data = json.load(stream)
    FKEYS = _frequency_pairs(data)
    for key, _ in FKEYS:
        if key not in data.get("results", {}):
            raise SystemExit(f"в results_interleaved.json нет варианта {key}")
        if "selected" not in data["results"][key]:
            raise SystemExit(f"в results_interleaved.json нет selected для {key}")
        if "quantities" not in data["results"][key]:
            raise SystemExit(f"в results_interleaved.json нет quantities для {key}")

    tn, fn = [1], [1]
    doc = setup_document()
    n_phases = data["tz"]["n_phases"]
    project = data.get("meta", {}).get("project", {})
    doc.core_properties.title = (
        f"Расчёт дросселя {n_phases}-фазного interleaved boost-преобразователя"
    )
    doc.core_properties.subject = (
        f"Предварительный аналитический расчёт: {_frequency_text()}"
    )
    doc.core_properties.author = project.get("name", "Расчётный проект")
    title_page(doc, data["tz"])
    sec_input(doc, data, tn)
    sec_mode(doc, data, tn, fn)
    sec_core(doc, data, tn)
    sec_magnetic(doc, data, tn)
    sec_winding(doc, data, tn, fn)
    sec_losses(doc, data, tn, fn)
    sec_parasitics(doc, data, tn, fn)
    sec_emc(doc, data, tn)
    sec_margins(doc, data, tn)
    sec_choice(doc, data, tn)
    sec_summary(doc, data, tn)
    sec_conclusions(doc, data)
    sec_scope(doc, data, tn)
    sec_sources(doc, data)

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"отчёт сохранён: {OUT_PATH}")
    print(f"таблиц: {tn[0] - 1}, рисунков: {fn[0] - 1}")


if __name__ == "__main__":
    main()
